import camelot
import pandas as pd
import re
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table, TableStyleInfo


STANDARD_COLUMNS = [
    "Transaction Date",
    "Value Date",
    "Description",
    "Reference Number",
    "Withdrawals",
    "Deposits",
    "Running Balance"
]


def is_sbi_bank_statement(bank_ledger):
    if bank_ledger is None:
        return False
    return "sbi" in str(bank_ledger).strip().lower()


# -----------------------------------------------------
# Header Utilities
# -----------------------------------------------------

def normalize_header(text):
    return (
        str(text)
        .lower()
        .replace("\n", " ")
        .replace("/", " ")
        .strip()
    )


def map_column(col):
    col_lower = normalize_header(col)

    if "transaction" in col_lower or ("txn" in col_lower and "date" in col_lower) or ("tran" in col_lower and "date" in col_lower):
        return "Transaction Date"

    if "value" in col_lower:
        return "Value Date"

    if "description" in col_lower or "particulars" in col_lower:
        return "Description"

    if "reference" in col_lower or "cheque" in col_lower or "ref" in col_lower or "chq" in col_lower:
        return "Reference Number"

    if "withdraw" in col_lower or "debit" in col_lower:
        return "Withdrawals"

    if "deposit" in col_lower or "credit" in col_lower:
        return "Deposits"

    if "balance" in col_lower:
        return "Running Balance"

    return None


# -----------------------------------------------------
# Repair Merged Deposit / Balance Columns
# -----------------------------------------------------

def repair_merged_amounts(df):

    pattern = re.compile(r"^\s*([\d,]+\.\d{2})\s+([\d,]+\.\d{2})\s*$")

    for i in range(len(df)):

        deposit = str(df.loc[i, "Deposits"]).strip()
        balance = str(df.loc[i, "Running Balance"]).strip()

        # Case 1: Deposits contains both values
        match = pattern.match(deposit)
        if match and balance == "":
            df.loc[i, "Deposits"] = match.group(1)
            df.loc[i, "Running Balance"] = match.group(2)
            continue

        # Case 2: Running Balance contains both values
        match = pattern.match(balance)
        if match and deposit == "":
            df.loc[i, "Deposits"] = match.group(1)
            df.loc[i, "Running Balance"] = match.group(2)

    return df


# -----------------------------------------------------
# Merge Spillover Rows (Multiline Transactions)
# -----------------------------------------------------

def merge_spillover_rows(df):

    rows_to_drop = []

    for i in range(1, len(df)):

        txn_date = str(df.loc[i, "Transaction Date"]).strip()
        val_date = str(df.loc[i, "Value Date"]).strip()
        withdrawal = str(df.loc[i, "Withdrawals"]).strip()
        deposit = str(df.loc[i, "Deposits"]).strip()
        desc = str(df.loc[i, "Description"]).strip()

        # Spillover condition:
        if (
            desc != "" and
            txn_date == "" and
            val_date == "" and
            withdrawal == "" and
            deposit == ""
        ):
            prev_desc = str(df.loc[i - 1, "Description"]).strip()
            df.loc[i - 1, "Description"] = f"{prev_desc} {desc}"
            rows_to_drop.append(i)

    df = df.drop(index=rows_to_drop).reset_index(drop=True)

    return df


def parse_amount(value):
    if pd.isna(value):
        return 0.0

    text = str(value).replace(",", "").strip()
    if text == "" or text.lower() in {"nan", "none"}:
        return 0.0

    try:
        return float(text)
    except ValueError:
        return 0.0


def clean_text(value):
    return re.sub(r"\s+", " ", str(value).replace("\n", " ")).strip()


def split_value_date_and_suffix(value):
    text = clean_text(value)
    if text == "":
        return None, None

    patterns = [
        r"^(\d{1,2}\s+[A-Za-z]{3,9}\s+20\d{2})\s+(.+)$",
        r"^(\d{1,2}[-/][A-Za-z]{3,9}[-/]20\d{2})\s+(.+)$",
        r"^(\d{1,2}[-/]\d{1,2}[-/]20\d{2})\s+(.+)$",
    ]

    for pattern in patterns:
        match = re.match(pattern, text, flags=re.IGNORECASE)
        if not match:
            continue

        parsed_date = pd.to_datetime(match.group(1), errors="coerce", dayfirst=True)
        if pd.isna(parsed_date):
            continue

        normalized_date = f"{parsed_date.day} {parsed_date.strftime('%b')} {parsed_date.year}"
        suffix = clean_text(match.group(2))
        if suffix == "":
            return None, None

        return normalized_date, suffix

    return None, None


def move_value_date_suffix_to_description(df):
    if df.empty:
        return df
    if "Value Date" not in df.columns or "Description" not in df.columns:
        return df

    for i in range(len(df)):
        value_date = clean_text(df.loc[i, "Value Date"])
        description = clean_text(df.loc[i, "Description"])

        if value_date != "":
            date_only, suffix = split_value_date_and_suffix(value_date)
            if not date_only or not suffix:
                continue

            if description.lower().startswith(suffix.lower()):
                merged_description = description
            else:
                merged_description = clean_text(f"{suffix} {description}")

            df.loc[i, "Value Date"] = date_only
            df.loc[i, "Description"] = merged_description
            continue

        # Inverse case: Value Date is empty but a leading date is merged into Description.
        desc_date, desc_remainder = split_value_date_and_suffix(description)
        if desc_date and desc_remainder:
            df.loc[i, "Value Date"] = desc_date
            df.loc[i, "Description"] = desc_remainder

    return df


def is_header_line(line):
    normalized = clean_text(line).upper()
    return (
        "TXN DATE" in normalized and
        "VALUE DATE" in normalized and
        "DESCRIPTION" in normalized and
        ("DEBITS" in normalized or "WITHDRAWALS" in normalized) and
        ("CREDITS" in normalized or "DEPOSITS" in normalized) and
        "BALANCE" in normalized
    )


def split_description_reference(description):
    # Best-effort reference extraction from long numeric ids in description.
    match = re.search(r"\b\d{6,}\b", description)
    reference = match.group(0) if match else ""
    return description, reference


def parse_text_transaction_line(line):
    date_regex = r"\d{2}-[A-Z]{3}-\d{4}"
    amount_regex = r"-?\d[\d,]*\.\d{2}"

    pattern = re.compile(
        rf"^({date_regex})\s+({date_regex})\s+(.*?)\s+({amount_regex})\s+({amount_regex})\s+({amount_regex})\s*$",
        re.IGNORECASE
    )

    match = pattern.match(clean_text(line))
    if not match:
        return None

    description, reference = split_description_reference(clean_text(match.group(3)))

    return {
        "Transaction Date": match.group(1).upper(),
        "Value Date": match.group(2).upper(),
        "Description": description,
        "Reference Number": reference,
        "Withdrawals": parse_amount(match.group(4)),
        "Deposits": parse_amount(match.group(5)),
        "Running Balance": parse_amount(match.group(6))
    }


def is_dbs_bank_statement(bank_ledger):
    if bank_ledger is None:
        return False
    return "dbs" in str(bank_ledger).strip().lower()


def is_kotak_bank_statement(bank_ledger):
    if bank_ledger is None:
        return False
    return "kotak" in str(bank_ledger).strip().lower()


def parse_dbs_amount(value):
    return parse_amount(value)


def amounts_equal(left, right):
    return round(abs(left - right), 2) == 0


def infer_dbs_amount_side(amount, balance, previous_balance):
    if previous_balance is None:
        return "", ""

    if amounts_equal(previous_balance - amount, balance):
        return amount, ""

    if amounts_equal(previous_balance + amount, balance):
        return "", amount

    return amount, ""


def parse_dbs_transaction_start(line, previous_balance):
    date_regex = r"\d{2}-[A-Za-z]{3}-\d{4}"
    amount_regex = r"\d{1,3}(?:,\d{2,3})*\.\d{2}|\d+\.\d{2}"
    pattern = re.compile(
        rf"^({date_regex})\s+({date_regex})\s+(.+?)\s+({amount_regex})(?:\s+({amount_regex}))?\s*$",
        re.IGNORECASE
    )

    match = pattern.match(clean_text(line))
    if not match:
        return None

    transaction_date = match.group(1)
    value_date = match.group(2)
    description = clean_text(match.group(3))
    first_amount = parse_dbs_amount(match.group(4))
    second_amount = parse_dbs_amount(match.group(5)) if match.group(5) else 0.0

    if second_amount > 0:
        balance = second_amount
        withdrawal, deposit = infer_dbs_amount_side(first_amount, balance, previous_balance)
    else:
        balance = first_amount
        withdrawal, deposit = "", ""

    description, reference = split_description_reference(description)

    return {
        "Transaction Date": transaction_date,
        "Value Date": value_date,
        "Description": description,
        "Reference Number": reference,
        "Withdrawals": withdrawal,
        "Deposits": deposit,
        "Running Balance": balance,
    }


def extract_dbs_treasures_statement(pdf_path):
    try:
        import pdfplumber
    except ImportError:
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    transactions = []
    current_transaction = None
    pending_description = []

    date_regex = re.compile(r"^\s*(\d{2}-\d{2}-\d{4})\s*")
    amount_regex = re.compile(r"([\d,]+\.\d{2})")

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text(layout=True) or ""
            for raw_line in text.splitlines():
                if not raw_line.strip():
                    continue
                
                lower_line = raw_line.lower()
                if "transaction history" in lower_line or "details of transaction" in lower_line:
                    continue
                if "private & confidential" in lower_line or "summary of account" in lower_line:
                    continue
                if "account alias" in lower_line or "saving bank a/c" in lower_line:
                    continue
                if "***end of transaction history***" in lower_line:
                    break
                if "dbs bank ltd. ground floor" in lower_line:
                    continue

                date_match = date_regex.match(raw_line)
                
                if date_match:
                    if current_transaction:
                        if pending_description:
                            current_transaction["Description"] += " " + " ".join(pending_description)
                            pending_description = []
                        current_transaction["Description"] = clean_text(current_transaction["Description"])
                        transactions.append(current_transaction)
                    else:
                        pending_description = []
                    
                    txn_date = date_match.group(1)
                    rest_of_line = raw_line[date_match.end():]
                    
                    withdrawal = ""
                    deposit = ""
                    
                    amounts = [(m.start(), m.group(1)) for m in amount_regex.finditer(raw_line)]
                    
                    if amounts:
                        last_amount_idx, last_amount_str = amounts[-1]
                        amount = parse_amount(last_amount_str)
                        if last_amount_idx < 68: # Debit
                            withdrawal = amount
                        else: # Credit
                            deposit = amount
                            
                        desc_part = raw_line[date_match.end():last_amount_idx].strip()
                    else:
                        desc_part = rest_of_line.strip()
                        
                    desc_parts = pending_description + [desc_part] if desc_part else pending_description
                    pending_description = []
                    
                    current_transaction = {
                        "Transaction Date": txn_date,
                        "Value Date": txn_date,
                        "Description": " ".join(desc_parts),
                        "Reference Number": "",
                        "Withdrawals": withdrawal,
                        "Deposits": deposit,
                        "Running Balance": ""
                    }
                    continue
                    
                desc = raw_line.strip()
                if current_transaction:
                    current_transaction["Description"] += " " + desc
                else:
                    pending_description.append(desc)
                    
    if current_transaction:
        if pending_description:
            current_transaction["Description"] += " " + " ".join(pending_description)
        current_transaction["Description"] = clean_text(current_transaction["Description"])
        transactions.append(current_transaction)
        
    if not transactions:
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    df = pd.DataFrame(transactions)
    for column in STANDARD_COLUMNS:
        if column not in df.columns:
            df[column] = ""

    df["Description"] = df["Description"].apply(clean_text)
    return df[STANDARD_COLUMNS]


def extract_dbs_new_format_statement(pdf_path):
    try:
        import pdfplumber
    except ImportError:
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    transactions = []
    current_transaction = None
    pending_description = []

    date_regex = re.compile(r"^\s*(\d{2}/\d{2}/\d{4})\s+(\d{2}/\d{2}/\d{4})\s+(\S+)\s+")
    amount_regex = re.compile(r"([\d,]+\.\d{2})")

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text(layout=True) or ""
            for raw_line in text.splitlines():
                if not raw_line.strip():
                    continue
                
                lower_line = raw_line.lower()
                if "transaction date" in lower_line or "account statement" in lower_line:
                    continue
                if "account details" in lower_line or "account number" in lower_line:
                    continue
                if "cheque/reference" in lower_line or "number" == lower_line.strip():
                    continue
                if "statement of account" in lower_line or "page " in lower_line:
                    continue
                if "dbs bank india ltd." in lower_line:
                    continue
                if "summary" in lower_line and "opening balance" in lower_line:
                    continue

                date_match = date_regex.match(raw_line)
                
                if date_match:
                    if current_transaction:
                        current_transaction["Description"] += " " + " ".join(pending_description)
                        current_transaction["Description"] = clean_text(current_transaction["Description"])
                        transactions.append(current_transaction)
                        
                    pending_description = []
                    
                    txn_date = date_match.group(1).replace("/", "-")
                    val_date = date_match.group(2).replace("/", "-")
                    branch = date_match.group(3)
                    
                    balance = 0.0
                    amount = 0.0
                    amt_idx = 0
                    
                    amounts = [(m.start(), m.group(1)) for m in amount_regex.finditer(raw_line)]
                    
                    if len(amounts) >= 2:
                        bal_idx, bal_str = amounts[-1]
                        amt_idx, amt_str = amounts[-2]
                        
                        balance = parse_amount(bal_str)
                        amount = parse_amount(amt_str)
                        desc_part = raw_line[date_match.end():amt_idx].strip()
                    else:
                        desc_part = raw_line[date_match.end():].strip()
                        
                    desc_parts = [desc_part] if desc_part else []
                    
                    current_transaction = {
                        "Transaction Date": txn_date,
                        "Value Date": val_date,
                        "Description": " ".join(desc_parts),
                        "Reference Number": "",
                        "_amount": amount,
                        "_balance": balance,
                        "_amt_idx": amt_idx
                    }
                    continue
                    
                desc = raw_line.strip()
                if current_transaction:
                    pending_description.append(desc)
                    
    if current_transaction:
        current_transaction["Description"] += " " + " ".join(pending_description)
        current_transaction["Description"] = clean_text(current_transaction["Description"])
        transactions.append(current_transaction)
        
    if not transactions:
        return pd.DataFrame(columns=STANDARD_COLUMNS)
        
    previous_balance = None
    final_transactions = []
    
    for txn in transactions:
        amount = txn.pop("_amount")
        balance = txn.pop("_balance")
        amt_idx = txn.pop("_amt_idx")
        
        withdrawal = ""
        deposit = ""
        
        if previous_balance is None:
            if amt_idx < 63:
                withdrawal = amount
            else:
                deposit = amount
        else:
            if round(abs(previous_balance - amount - balance), 2) == 0:
                withdrawal = amount
            elif round(abs(previous_balance + amount - balance), 2) == 0:
                deposit = amount
            else:
                if amt_idx < 63:
                    withdrawal = amount
                else:
                    deposit = amount
                    
        previous_balance = balance
        txn["Withdrawals"] = withdrawal
        txn["Deposits"] = deposit
        txn["Running Balance"] = balance
        final_transactions.append(txn)
        
    df = pd.DataFrame(final_transactions)
    for column in STANDARD_COLUMNS:
        if column not in df.columns:
            df[column] = ""

    df["Description"] = df["Description"].apply(clean_text)
    return df[STANDARD_COLUMNS]


def extract_dbs_text_statement(pdf_path):
    try:
        import pdfplumber
    except ImportError:
        print("pdfplumber is not installed. DBS text fallback unavailable.")
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    with pdfplumber.open(pdf_path) as pdf:
        first_page = pdf.pages[0].extract_text() or ""
        if "Details of Transaction" in first_page and "Debit" in first_page:
            return extract_dbs_treasures_statement(pdf_path)
        if "Branch code" in first_page and "Balance" in first_page:
            return extract_dbs_new_format_statement(pdf_path)

    transactions = []
    current_transaction = None
    previous_balance = None
    in_savings_statement = False
    date_start_pattern = re.compile(r"^\d{2}-[A-Za-z]{3}-\d{4}\s+\d{2}-[A-Za-z]{3}-\d{4}\s+", re.IGNORECASE)
    opening_balance_pattern = re.compile(r"^Opening Balance\s+([\d,]+\.\d{2})$", re.IGNORECASE)

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""

            for raw_line in text.splitlines():
                line = clean_text(raw_line)

                if line == "" or is_header_line(line):
                    continue

                if "Account Type: SAVINGS" in line:
                    in_savings_statement = True
                    continue

                if not in_savings_statement:
                    continue

                if line.startswith("Closing Balance"):
                    if current_transaction:
                        transactions.append(current_transaction)
                        current_transaction = None
                    in_savings_statement = False
                    continue

                opening_match = opening_balance_pattern.match(line)
                if opening_match:
                    previous_balance = parse_dbs_amount(opening_match.group(1))
                    continue

                if date_start_pattern.match(line):
                    parsed = parse_dbs_transaction_start(line, previous_balance)
                    if not parsed:
                        continue

                    if current_transaction:
                        transactions.append(current_transaction)

                    current_transaction = parsed
                    previous_balance = parse_dbs_amount(parsed["Running Balance"])
                    continue

                if current_transaction:
                    current_transaction["Description"] = clean_text(
                        f"{current_transaction['Description']} {line}"
                    )
                    current_transaction["Description"], current_transaction["Reference Number"] = split_description_reference(
                        current_transaction["Description"]
                    )

    if current_transaction:
        transactions.append(current_transaction)

    if not transactions:
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    df = pd.DataFrame(transactions)
    for column in STANDARD_COLUMNS:
        if column not in df.columns:
            df[column] = ""

    df["Description"] = df["Description"].apply(clean_text)
    return df[STANDARD_COLUMNS]


def parse_kotak_transaction_start(line, previous_balance):
    date_regex = r"\d{1,2}\s+[A-Za-z]{3}\s+20\d{2}"
    amount_regex = r"\d{1,3}(?:,\d{2,3})*\.\d{2}|\d+\.\d{2}"
    pattern = re.compile(
        rf"^\d+\s+({date_regex})\s+(.+?)\s+({amount_regex})\s+({amount_regex})\s*$",
        re.IGNORECASE
    )

    match = pattern.match(clean_text(line))
    if not match:
        return None

    transaction_date = match.group(1)
    description = clean_text(match.group(2))
    amount = parse_amount(match.group(3))
    balance = parse_amount(match.group(4))
    withdrawal, deposit = infer_dbs_amount_side(amount, balance, previous_balance)
    description, reference = split_description_reference(description)

    return {
        "Transaction Date": transaction_date,
        "Value Date": transaction_date,
        "Description": description,
        "Reference Number": reference,
        "Withdrawals": withdrawal,
        "Deposits": deposit,
        "Running Balance": balance,
    }


def extract_kotak_text_statement(pdf_path):
    try:
        import pdfplumber
    except ImportError:
        print("pdfplumber is not installed. Kotak text fallback unavailable.")
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    transactions = []
    current_transaction = None
    previous_balance = None
    in_transaction_section = False
    transaction_start_pattern = re.compile(r"^\d+\s+\d{1,2}\s+[A-Za-z]{3}\s+20\d{2}\s+", re.IGNORECASE)
    opening_balance_pattern = re.compile(r"^-\s+-\s+Opening Balance\s+-\s+-\s+-\s+([\d,]+\.\d{2})$", re.IGNORECASE)

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""

            for raw_line in text.splitlines():
                line = clean_text(raw_line)

                if line == "":
                    continue

                if "Current Account Transactions" in line:
                    in_transaction_section = True
                    continue

                if not in_transaction_section:
                    continue

                if (
                    line.startswith("# Date Description") or
                    line.startswith("Statement Generated") or
                    line.startswith("Account Statement") or
                    line.startswith("Account No.") or
                    line == "NSB TRADERS"
                ):
                    continue

                opening_match = opening_balance_pattern.match(line)
                if opening_match:
                    previous_balance = parse_amount(opening_match.group(1))
                    continue

                if transaction_start_pattern.match(line):
                    parsed = parse_kotak_transaction_start(line, previous_balance)
                    if not parsed:
                        continue

                    if current_transaction:
                        transactions.append(current_transaction)

                    current_transaction = parsed
                    previous_balance = parse_amount(parsed["Running Balance"])
                    continue

                if current_transaction:
                    current_transaction["Description"] = clean_text(
                        f"{current_transaction['Description']} {line}"
                    )
                    current_transaction["Description"], current_transaction["Reference Number"] = split_description_reference(
                        current_transaction["Description"]
                    )

    if current_transaction:
        transactions.append(current_transaction)

    if not transactions:
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    df = pd.DataFrame(transactions)
    for column in STANDARD_COLUMNS:
        if column not in df.columns:
            df[column] = ""

    df["Description"] = df["Description"].apply(clean_text)
    return df[STANDARD_COLUMNS]


def extract_text_statement(pdf_path):
    try:
        import pdfplumber
    except ImportError:
        print("pdfplumber is not installed. Text fallback unavailable.")
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    transactions = []
    current_transaction = None
    date_start_pattern = re.compile(r"^\d{2}-[A-Z]{3}-\d{4}", re.IGNORECASE)

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""

            for raw_line in text.splitlines():
                line = clean_text(raw_line)

                if line == "" or is_header_line(line):
                    continue

                if date_start_pattern.match(line):
                    parsed = parse_text_transaction_line(line)
                    if not parsed:
                        continue

                    if current_transaction:
                        transactions.append(current_transaction)

                    current_transaction = parsed
                elif current_transaction:
                    # Non-date lines are treated as spillover description lines.
                    current_transaction["Description"] = clean_text(
                        f"{current_transaction['Description']} {line}"
                    )

    if current_transaction:
        transactions.append(current_transaction)

    if not transactions:
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    df = pd.DataFrame(transactions)

    for column in STANDARD_COLUMNS:
        if column not in df.columns:
            df[column] = ""

    df["Description"] = df["Description"].apply(clean_text)
    df = df[STANDARD_COLUMNS]

    return df


def has_valid_transactions(df):
    required = {"Transaction Date", "Value Date", "Description", "Withdrawals", "Deposits"}
    if df.empty or not required.issubset(set(df.columns)):
        return False

    for _, row in df.iterrows():
        description = clean_text(row.get("Description", ""))
        withdrawal = parse_amount(row.get("Withdrawals"))
        deposit = parse_amount(row.get("Deposits"))

        txn_date = pd.to_datetime(row.get("Transaction Date"), errors="coerce")
        val_date = pd.to_datetime(row.get("Value Date"), errors="coerce")

        if description and (withdrawal > 0 or deposit > 0) and pd.notna(txn_date) and pd.notna(val_date):
            return True

    return False


def extract_table_statement(pdf_path):
    tables = camelot.read_pdf(pdf_path, pages="all")

    print(f"Total tables detected: {tables.n}")

    dataframes = []
    accepted_tables = 0
    ignored_tables = 0

    for table in tables:

        df = table.df

        if df.empty or df.shape[0] < 2:
            ignored_tables += 1
            continue

        # First row as header
        df.columns = df.iloc[0]
        df = df[1:]

        # Dynamic column mapping
        mapped_columns = {}
        for col in df.columns:
            mapped = map_column(col)
            if mapped:
                mapped_columns[col] = mapped

        df = df.rename(columns=mapped_columns)

        # Required fields
        required = {
            "Transaction Date",
            "Value Date",
            "Description",
            "Withdrawals",
            "Deposits"
        }

        if not required.issubset(set(df.columns)):
            ignored_tables += 1
            continue

        # Add optional columns if missing
        if "Reference Number" not in df.columns:
            df["Reference Number"] = ""

        if "Running Balance" not in df.columns:
            df["Running Balance"] = ""

        # Standardize column order
        df = df[STANDARD_COLUMNS]

        accepted_tables += 1
        dataframes.append(df)

    final_df = pd.DataFrame()
    if dataframes:
        final_df = pd.concat(dataframes, ignore_index=True)

        # Clean all cell values
        for col in final_df.columns:
            final_df[col] = (
                final_df[col]
                .astype(str)
                .str.replace("\n", " ", regex=False)
                .str.replace("\r", "", regex=False)
                .str.replace(r"\s+", " ", regex=True)
                .str.strip()
            )

        # Repair merged columns
        final_df = repair_merged_amounts(final_df)

        # Merge multiline spillovers
        final_df = merge_spillover_rows(final_df)

    return final_df, accepted_tables, ignored_tables


# -----------------------------------------------------
# Main Extraction Function
# -----------------------------------------------------

def is_icici_bank_statement(bank_ledger):
    if bank_ledger is None:
        return False
    return "icici" in str(bank_ledger).strip().lower()


def extract_icici_text_statement(pdf_path):
    try:
        import pdfplumber
    except ImportError:
        print("pdfplumber is not installed. ICICI text fallback unavailable.")
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    transactions = []
    current_transaction = None
    previous_balance = None
    pending_description = []

    date_regex = r"\d{2}-\d{2}-\d{4}"
    amount_regex = r"[\d,]+\.\d{2}"
    
    txn_pattern = re.compile(
        rf"^({date_regex})\s*({date_regex})\s*(.*?)\s*({amount_regex})\s+({amount_regex})(?:\s+(Cr|Dr))?$",
        re.IGNORECASE
    )
    
    bf_pattern = re.compile(
        rf"^({date_regex})\s*B/F\s*(.*?)\s*({amount_regex})(?:\s+(Cr|Dr))?$",
        re.IGNORECASE
    )

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            text = page.extract_text() or ""
            for raw_line in text.splitlines():
                line = clean_text(raw_line)
                if not line: continue
                
                bf_match = bf_pattern.match(line)
                if bf_match:
                    pending_description = []
                    previous_balance = parse_amount(bf_match.group(3))
                    if current_transaction:
                        transactions.append(current_transaction)
                        current_transaction = None
                    continue
                
                txn_match = txn_pattern.match(line)
                if txn_match:
                    if current_transaction:
                        if pending_description:
                            current_transaction["Description"] += " " + " ".join(pending_description)
                            pending_description = []
                        current_transaction["Description"] = clean_text(current_transaction["Description"])
                        transactions.append(current_transaction)
                    else:
                        pending_description = []
                    
                    txn_date = txn_match.group(1)
                    val_date = txn_match.group(2)
                    location = txn_match.group(3)
                    amount = parse_amount(txn_match.group(4))
                    balance = parse_amount(txn_match.group(5))
                    
                    withdrawal, deposit = infer_dbs_amount_side(amount, balance, previous_balance)
                    previous_balance = balance
                    
                    desc_parts = pending_description + [location]
                    pending_description = []
                    
                    current_transaction = {
                        "Transaction Date": txn_date,
                        "Value Date": val_date,
                        "Description": " ".join(desc_parts),
                        "Reference Number": "",
                        "Withdrawals": withdrawal,
                        "Deposits": deposit,
                        "Running Balance": balance
                    }
                    continue
                
                lower_line = line.lower()
                if "operative account in inr" in lower_line or "statement of transactions" in lower_line or "tran date value date" in lower_line:
                    continue
                if "page " in lower_line and " of " in lower_line:
                    continue
                if line.startswith("Total :") or line.startswith("Summary of Accounts") or line.startswith("Your Base Branch"):
                    continue
                if "your details with us" in lower_line or "mr." in lower_line or "plot no" in lower_line or "hyderabad" in lower_line or "telangana" in lower_line:
                    continue
                    
                if current_transaction:
                    current_transaction["Description"] += " " + line
                else:
                    pending_description.append(line)
                    
    if current_transaction:
        if pending_description:
            current_transaction["Description"] += " " + " ".join(pending_description)
        current_transaction["Description"] = clean_text(current_transaction["Description"])
        transactions.append(current_transaction)

    if not transactions:
        return pd.DataFrame(columns=STANDARD_COLUMNS)

    df = pd.DataFrame(transactions)
    for column in STANDARD_COLUMNS:
        if column not in df.columns:
            df[column] = ""

    df["Description"] = df["Description"].apply(clean_text)
    return df[STANDARD_COLUMNS]


def extract_bank_statement(pdf_path, output_file=None, bank_ledger=None):
    final_df = pd.DataFrame()
    accepted_tables = 0
    ignored_tables = 0
    extraction_method = "table"

    sbi_statement = is_sbi_bank_statement(bank_ledger)
    dbs_statement = is_dbs_bank_statement(bank_ledger)
    kotak_statement = is_kotak_bank_statement(bank_ledger)
    icici_statement = is_icici_bank_statement(bank_ledger)

    if sbi_statement:
        print("SBI detected: using text-based row extraction by default...")
        final_df = extract_text_statement(pdf_path)
        final_df = move_value_date_suffix_to_description(final_df)
        extraction_method = "sbi_text"

    if dbs_statement:
        print("DBS detected: using text-table row extraction by default...")
        final_df = extract_dbs_text_statement(pdf_path)
        extraction_method = "dbs_text_table"

    if kotak_statement:
        print("Kotak detected: using text-based row extraction by default...")
        final_df = extract_kotak_text_statement(pdf_path)
        extraction_method = "kotak_text"

    if icici_statement:
        print("ICICI detected: using text-based row extraction by default...")
        final_df = extract_icici_text_statement(pdf_path)
        extraction_method = "icici_text"

    if not has_valid_transactions(final_df):
        final_df, accepted_tables, ignored_tables = extract_table_statement(pdf_path)
        final_df = move_value_date_suffix_to_description(final_df)
        extraction_method = "table"

    if not has_valid_transactions(final_df):
        if not any([sbi_statement, dbs_statement, kotak_statement, icici_statement]):
            print("Switching to text-based extraction (pdfplumber fallback)...")
            final_df = extract_text_statement(pdf_path)
            final_df = move_value_date_suffix_to_description(final_df)
            extraction_method = "text"

    if not has_valid_transactions(final_df):
        raise ValueError("No valid tables found in PDF. Extraction aborted.")

    if output_file:
        # Save Excel
        final_df.to_excel(output_file, index=False)

        # Format as Excel table
        wb = load_workbook(output_file)
        ws = wb.active

        last_row = ws.max_row
        last_col = ws.max_column
        table_range = f"A1:{ws.cell(row=last_row, column=last_col).coordinate}"

        excel_table = Table(displayName="CombinedTable", ref=table_range)

        style = TableStyleInfo(
            name="TableStyleMedium9",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False
        )

        excel_table.tableStyleInfo = style
        ws.add_table(excel_table)

        wb.save(output_file)

    if extraction_method in {"dbs_text_table", "kotak_text", "icici_text"}:
        if extraction_method == "dbs_text_table":
            print("Extraction method: DBS text-table parser")
        elif extraction_method == "icici_text":
            print("Extraction method: ICICI text parser")
        else:
            print("Extraction method: Kotak text parser")
        print(f"Transaction rows extracted: {len(final_df)}")
    else:
        print(f"Tables accepted: {accepted_tables}")
        print(f"Tables ignored : {ignored_tables}")

    print("Bank statement extraction completed.")

    return final_df
