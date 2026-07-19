import os
import sys
import pandas as pd
from openpyxl import load_workbook

from core.duplicate_filter import ExistingTransactionMatcher, JSON_DUPLICATE_COLUMNS
from core.client_config import (
    extract_bank_code,
    get_default_client_name,
    is_import_enabled,
    resolve_duplicate_json_path,
    resolve_rule_path,
)
from core.rule_engine import RuleEngine
from utils.file_writer import safe_excel_write


if len(sys.argv) > 1:
    PURCHASES_FILE_PATH = sys.argv[1]
else:
    PURCHASES_FILE_PATH = "./input/Purchases/Arjunrao Jan26 Purchase.xlsx"

if len(sys.argv) > 2:
    BANK_LEDGER = sys.argv[2]
else:
    BANK_LEDGER = "494"

if len(sys.argv) > 3:
    CLIENT_NAME = sys.argv[3]
else:
    CLIENT_NAME = get_default_client_name()

FINAL_OUTPUT = "./output/Purchases_Import.xlsx"
DUPLICATE_JSON_PATH = resolve_duplicate_json_path(CLIENT_NAME)

REQUIRED_COLUMNS = [
    "NAME",
    "INVOICE NO",
    "DATE",
    "GROSS AMT",
]

COLUMN_ALIASES = {
    "NAME": "NAME",
    "PARTICULARS": "NAME",
    "PARTICULAR": "NAME",
    "SUPPLIER": "NAME",
    "SUPPLIER NAME": "NAME",
    "PARTY": "NAME",
    "PARTY NAME": "NAME",
    "INVOICE NO": "INVOICE NO",
    "INVOICE NO_": "INVOICE NO",
    "INVOICE NUMBER": "INVOICE NO",
    "BILLNO": "INVOICE NO",
    "BILL NO": "INVOICE NO",
    "BILL NUMBER": "INVOICE NO",
    "DATE": "DATE",
    "BILL DATE": "DATE",
    "INVOICE DATE": "DATE",
    "GROSS AMT": "GROSS AMT",
    "GROSS AMOUNT": "GROSS AMT",
    "INVOICE VALUE": "GROSS AMT",
    "BILL VALUE": "GROSS AMT",
    "AMOUNT": "GROSS AMT",
}

IMPORT_COLUMNS = [
    "Voucher_Num",
    "Voucher_Type",
    "Date",
    "Description",
    "Narration",
    "Cr_Ledger",
    "Amount",
    "Cr",
    "Dr_Ledger",
    "Dr_Amount",
    "Dr",
]

UNCLASSIFIED_COLUMNS = [
    "Unclassified_Num",
    "Name",
    "Invoice No.",
    "Date",
    "Gross AMT",
    "Reason",
]


class _PurchaseTransaction:
    def __init__(self, description):
        self.description = description


def _normalize_columns(columns):
    return {col: _canonical_column_name(col) for col in columns}


def _normalize_header_cell(value):
    if pd.isna(value):
        return ""
    return " ".join(str(value).strip().upper().replace(".", "").split())


def _canonical_column_name(value):
    normalized = _normalize_header_cell(value)
    return COLUMN_ALIASES.get(normalized, normalized)


def _dataframe_from_detected_header(raw_df):
    if raw_df.empty:
        return pd.DataFrame()

    required_set = set(REQUIRED_COLUMNS)
    best_row = None
    best_score = -1
    scan_limit = min(len(raw_df), 25)

    for idx in range(scan_limit):
        row_values = raw_df.iloc[idx].tolist()
        normalized = [_canonical_column_name(v) for v in row_values]
        score = len(required_set.intersection(normalized))
        if score > best_score:
            best_score = score
            best_row = idx
        if score == len(required_set):
            break

    if best_row is None or best_score <= 0:
        return raw_df

    headers = []
    seen = {}
    for value in raw_df.iloc[best_row].tolist():
        header = str(value).strip() if not pd.isna(value) else ""
        if header in seen:
            seen[header] += 1
            header = f"{header}_{seen[header]}"
        else:
            seen[header] = 0
        headers.append(header)

    body = raw_df.iloc[best_row + 1 :].copy()
    body.columns = headers
    body = body.dropna(how="all")
    return body.reset_index(drop=True)


def _parse_amount(value):
    if pd.isna(value):
        return 0.0

    text = str(value).replace(",", "").strip()
    if text == "":
        return 0.0

    try:
        return float(text)
    except ValueError:
        return 0.0


def _format_date(value):
    parsed = pd.to_datetime(value, errors="coerce", dayfirst=True)
    if pd.isna(parsed):
        return ""
    return parsed.strftime("%d-%m-%Y")


def _read_worksheet_table(ws):
    if ws.tables:
        first_table_name = next(iter(ws.tables))
        table = ws.tables[first_table_name]

        table_data = [
            [cell.value for cell in row]
            for row in ws[table.ref]
        ]

        if not table_data:
            return pd.DataFrame()

        headers = [
            str(col).strip() if col is not None else ""
            for col in table_data[0]
        ]
        rows = table_data[1:]
        return pd.DataFrame(rows, columns=headers)

    return None


def _read_purchases_table(purchases_file_path):
    wb = load_workbook(purchases_file_path, data_only=True)
    sheet_frames = []

    for ws in wb.worksheets:
        table_df = _read_worksheet_table(ws)
        if table_df is not None:
            sheet_df = table_df
        else:
            raw_df = pd.read_excel(purchases_file_path, sheet_name=ws.title, header=None)
            sheet_df = _dataframe_from_detected_header(raw_df)

        if not sheet_df.empty:
            sheet_frames.append(sheet_df)

    if not sheet_frames:
        return pd.DataFrame()

    return pd.concat(sheet_frames, ignore_index=True, sort=False)


def _coalesce_duplicate_columns(df):
    if not df.columns.has_duplicates:
        return df

    output = pd.DataFrame(index=df.index)
    for column in dict.fromkeys(df.columns):
        same_name = df.loc[:, df.columns == column]
        if same_name.shape[1] == 1:
            output[column] = same_name.iloc[:, 0]
        else:
            output[column] = same_name.bfill(axis=1).iloc[:, 0]
    return output


def _resolve_party_ledger(name, rule_engine):
    txn = _PurchaseTransaction(name)
    rule = rule_engine.match(txn)
    if rule and rule.get("ledger"):
        return rule["ledger"], True
    return "", False


def _empty_import_df():
    return pd.DataFrame(columns=IMPORT_COLUMNS)


def _empty_unclassified_df():
    return pd.DataFrame(columns=UNCLASSIFIED_COLUMNS)


def _empty_duplicate_df():
    return pd.DataFrame(columns=["Duplicate_Num"] + IMPORT_COLUMNS[1:] + JSON_DUPLICATE_COLUMNS)


def _prepare_import_df(rows):
    if not rows:
        return _empty_import_df()

    output_df = pd.DataFrame(rows)
    output_df.reset_index(drop=True, inplace=True)
    output_df.insert(0, "Voucher_Num", output_df.index + 1)
    return output_df.reindex(columns=IMPORT_COLUMNS)


def _prepare_unclassified_df(rows):
    if not rows:
        return _empty_unclassified_df()

    output_df = pd.DataFrame(rows)
    output_df.reset_index(drop=True, inplace=True)
    output_df.insert(0, "Unclassified_Num", output_df.index + 1)
    return output_df.reindex(columns=UNCLASSIFIED_COLUMNS)


def _prepare_duplicate_df(rows):
    if not rows:
        return _empty_duplicate_df()

    output_df = pd.DataFrame(rows)
    output_df.reset_index(drop=True, inplace=True)
    output_df.insert(0, "Duplicate_Num", output_df.index + 1)
    columns = ["Duplicate_Num"] + IMPORT_COLUMNS[1:] + JSON_DUPLICATE_COLUMNS
    return output_df.reindex(columns=columns)


def build_purchase_sheets(purchases_file_path, bank_ledger, client_name):
    df = _read_purchases_table(purchases_file_path)
    df = df.rename(columns=_normalize_columns(df.columns))
    df = _coalesce_duplicate_columns(df)

    missing = [col for col in REQUIRED_COLUMNS if col not in df.columns]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    active_rule_path = resolve_rule_path(bank_ledger, client_name)
    rule_engine = RuleEngine(active_rule_path)

    duplicate_matcher = None
    if os.path.exists(DUPLICATE_JSON_PATH):
        duplicate_matcher = ExistingTransactionMatcher(DUPLICATE_JSON_PATH)

    import_rows = []
    unclassified_rows = []
    duplicate_rows = []

    for _, row in df.iterrows():
        date = _format_date(row.get("DATE"))
        name = str(row.get("NAME", "")).strip()
        amount = _parse_amount(row.get("GROSS AMT"))
        invoice_no = str(row.get("INVOICE NO", "")).strip()

        if not date or name == "" or amount <= 0:
            continue

        party_ledger, matched = _resolve_party_ledger(name, rule_engine)
        if not matched:
            unclassified_rows.append(
                {
                    "Name": name,
                    "Invoice No.": invoice_no,
                    "Date": date,
                    "Gross AMT": amount,
                    "Reason": "No matching description rule",
                }
            )
            continue

        voucher_row = {
            "Voucher_Type": "Journal",
            "Date": date,
            "Description": name,
            "Narration": invoice_no,
            "Cr_Ledger": party_ledger,
            "Amount": amount,
            "Cr": "CR",
            "Dr_Ledger": "Purchases",
            "Dr_Amount": amount,
            "Dr": "DR",
        }

        if duplicate_matcher:
            matched_record = duplicate_matcher.find_voucher_match(
                voucher_row.get("Date"),
                voucher_row.get("Amount"),
                [voucher_row.get("Cr_Ledger"), voucher_row.get("Dr_Ledger")],
            )
            if matched_record:
                duplicate_row = dict(voucher_row)
                duplicate_row.update(matched_record["json_details"])
                duplicate_rows.append(duplicate_row)
                continue

        import_rows.append(voucher_row)

    return (
        _prepare_import_df(import_rows),
        _prepare_unclassified_df(unclassified_rows),
        _prepare_duplicate_df(duplicate_rows),
        active_rule_path,
    )


def main():
    if not is_import_enabled(CLIENT_NAME, "Purchases", BANK_LEDGER):
        print(f"\nPurchases module is not configured for bank '{extract_bank_code(BANK_LEDGER)}' under {CLIENT_NAME}.")
        return

    import_df, unclassified_df, duplicate_df, rule_path = build_purchase_sheets(
        PURCHASES_FILE_PATH,
        BANK_LEDGER,
        CLIENT_NAME,
    )

    def write_workbook():
        with pd.ExcelWriter(FINAL_OUTPUT, engine="openpyxl") as writer:
            import_df.to_excel(writer, sheet_name="Import", index=False)
            unclassified_df.to_excel(writer, sheet_name="Unclassified", index=False)
            duplicate_df.to_excel(writer, sheet_name="Duplicate", index=False)

    safe_excel_write(write_workbook, FINAL_OUTPUT)

    print("\nPurchases import workbook generated successfully.")
    print(f"Client used       : {CLIENT_NAME}")
    print(f"Bank used         : {BANK_LEDGER}")
    print(f"Rules file used   : {rule_path}")
    print(f"Imported rows     : {len(import_df)}")
    print(f"Unclassified rows : {len(unclassified_df)}")
    print(f"Duplicate rows    : {len(duplicate_df)}")
    print(f"Output file       : {FINAL_OUTPUT}")


if __name__ == "__main__":
    main()
